from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
FILES_DIR = (BASE_DIR / "files").resolve()


def list_excel_files():
    """Faqat files papkasidagi Excel fayllarni topadi."""
    FILES_DIR.mkdir(parents=True, exist_ok=True)

    return sorted(
        p for p in FILES_DIR.iterdir()
        if p.is_file()
        and not p.name.startswith("~$")
        and p.suffix.lower() in {".xlsx", ".xlsm"}
    )


def inspect_excel(filename):
    """Excel faylining varaqlari va to'ldirilgan kataklarini ko'rsatadi."""
    allowed_files = {p.name: p for p in list_excel_files()}

    if filename not in allowed_files:
        raise ValueError("Fayl ruxsat berilgan papkada topilmadi.")

    path = allowed_files[filename]
    keep_vba = path.suffix.lower() == ".xlsm"

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=keep_vba
    )

    result = []
    try:
        for sheet in workbook.worksheets:
            result.append(f"Varaq: {sheet.title}")
            count = 0

            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        result.append(
                            f"{cell.coordinate} = {cell.value!r}"
                        )
                        count += 1

                        if count >= 100:
                            result.append(
                                "... ko'rsatish limiti: 100 ta katak"
                            )
                            break

                if count >= 100:
                    break

            if count == 0:
                result.append("(Bo'sh varaq)")

    finally:
        workbook.close()

    return "\n".join(result)


if __name__ == "__main__":
    files = list_excel_files()

    if not files:
        print("Excel fayllar topilmadi.")
        print(f"Fayllarni shu papkaga joylashtiring: {FILES_DIR}")
    else:
        print("Topilgan Excel fayllar:")
        for file in files:
            print(f"- {file.name}")

        print("\nFayl tarkibini ko'rish uchun fayl nomini kiriting.")
        filename = input("> ").strip()

        try:
            print(inspect_excel(filename))
        except Exception as error:
            print(f"Xatolik: {error}")
